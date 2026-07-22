from fastapi.testclient import TestClient
from io import BytesIO
import base64

from docx import Document
from docx.shared import Inches

from app.main import app
from app.services import document_service
from app.services.ai_service import AIServiceError

import pytest


client = TestClient(app)


@pytest.fixture(autouse=True)
def stub_ai(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(
        document_service,
        "improve_article",
        lambda text, filename, connection: f"# AI 整理结果\n\n{text}",
    )
    monkeypatch.setattr(
        document_service,
        "upload_task_assets",
        lambda task_id, assets: assets,
    )


def test_health() -> None:
    response = client.get("/api/health")
    assert response.status_code == 200
    assert response.json()["status"] == "ok"


def test_markdown_workflow() -> None:
    content = "# 测试文章\n\n联系我：13812345678\n"
    response = client.post(
        "/api/documents/analyze",
        files={"file": ("draft.md", content.encode("utf-8"), "text/markdown")},
    )
    assert response.status_code == 200
    payload = response.json()
    assert payload["source_type"] == "markdown"
    assert payload["stats"]["findings"] == 1

    finding_id = payload["findings"][0]["id"]
    processed = client.post(
        f"/api/documents/{payload['id']}/process",
        json={"finding_ids": [finding_id], "improve_structure": True},
    )
    assert processed.status_code == 200
    assert "138****5678" in processed.json()["processed_content"]
    assert processed.json()["processed_content"].startswith("# AI 整理结果")

    for format_name in ("markdown", "hexo", "hugo", "docx", "pdf"):
        exported = client.get(f"/api/documents/{payload['id']}/export/{format_name}")
        assert exported.status_code == 200
        assert exported.content


def test_rejects_unsupported_file() -> None:
    response = client.post(
        "/api/documents/analyze",
        files={"file": ("draft.txt", b"hello", "text/plain")},
    )
    assert response.status_code == 400


def test_docx_image_is_added_to_published_content(monkeypatch: pytest.MonkeyPatch) -> None:
    png = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
    )
    document = Document()
    document.add_paragraph("图片测试")
    document.add_picture(BytesIO(png), width=Inches(1))
    output = BytesIO()
    document.save(output)

    def uploaded(task_id: str, assets: list[dict]) -> list[dict]:
        for asset in assets:
            asset.update(status="uploaded", url="https://cdn.example.com/test.png", provider="picgo")
        return assets

    monkeypatch.setattr(document_service, "upload_task_assets", uploaded)
    analyzed = client.post(
        "/api/documents/analyze",
        files={"file": ("images.docx", output.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert analyzed.status_code == 200
    payload = analyzed.json()
    assert len(payload["assets"]) == 1
    preview = client.get(
        f"/api/documents/{payload['id']}/assets/{payload['assets'][0]['filename']}"
    )
    assert preview.status_code == 200
    assert preview.content == png
    processed = client.post(
        f"/api/documents/{payload['id']}/process",
        json={"finding_ids": [], "improve_structure": False},
    )
    assert processed.status_code == 200
    assert "https://cdn.example.com/test.png" in processed.json()["processed_content"]
    assert processed.json()["assets"][0]["provider"] == "picgo"


def test_ai_failure_returns_specific_error(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(
        document_service,
        "improve_article",
        lambda *args, **kwargs: (_ for _ in ()).throw(AIServiceError("服务超时")),
    )
    analyzed = client.post(
        "/api/documents/analyze",
        files={"file": ("fallback.md", "没有标题的正文".encode(), "text/markdown")},
    ).json()
    response = client.post(
        f"/api/documents/{analyzed['id']}/process",
        json={"finding_ids": [], "improve_structure": True},
    )
    assert response.status_code == 502
    assert response.json()["detail"] == "服务超时"
