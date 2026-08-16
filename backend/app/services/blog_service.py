from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass, field
from datetime import datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path
from typing import Literal

import fitz
import yaml
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


class BlockType(str):
    HEADING = "heading"
    PARAGRAPH = "paragraph"
    IMAGE = "image"
    TABLE = "table"
    CODE = "code"
    LIST = "list"


@dataclass
class ImageBlock:
    index: int
    data: bytes
    ext: str
    md5: str
    paragraph_idx: int
    alt_text: str = ""


@dataclass
class ContentBlock:
    type: str
    content: str = ""
    level: int = 0
    language: str = ""
    image: ImageBlock | None = None


@dataclass
class ParsedDocument:
    blocks: list[ContentBlock] = field(default_factory=list)
    images: list[ImageBlock] = field(default_factory=list)
    detected_lang: str = "zh"
    title_hint: str = ""


def parse_docx(content: bytes) -> ParsedDocument:
    try:
        document = Document(BytesIO(content))
    except Exception as exc:
        raise ValueError("DOCX 文件无法读取，可能已经损坏或加密。") from exc

    blocks: list[ContentBlock] = []
    images: list[ImageBlock] = []
    seen_md5: dict[str, int] = {}
    para_idx = 0

    for element in document.element.body:
        tag = element.tag.split("}")[-1] if "}" in element.tag else element.tag
        if tag == "p":
            paragraph = Paragraph(element, document)
            block, image = _parse_docx_paragraph(paragraph, para_idx, seen_md5, document)
            if image:
                images.append(image)
            if block:
                blocks.append(block)
            para_idx += 1
        elif tag == "tbl":
            table_block = _parse_docx_table(Table(element, document))
            if table_block:
                blocks.append(table_block)

    return ParsedDocument(
        blocks=blocks,
        images=images,
        detected_lang=_detect_language(" ".join(p.text for p in document.paragraphs[:20])),
        title_hint=_extract_docx_title(document),
    )


def parse_pdf(content: bytes) -> ParsedDocument:
    try:
        document = fitz.open(stream=content, filetype="pdf")
    except Exception as exc:
        raise ValueError("PDF 文件无法读取，可能已经损坏或加密。") from exc
    if document.needs_pass:
        raise ValueError("暂不支持加密 PDF，请先解除密码。")

    blocks: list[ContentBlock] = []
    images: list[ImageBlock] = []
    seen_md5: dict[str, int] = {}

    try:
        for page in document:
            page_dict = page.get_text("dict")
            for raw_block in page_dict.get("blocks", []):
                if raw_block.get("type") == 0:
                    text, size, is_bold, is_mono = _extract_pdf_text_block(raw_block)
                    if text:
                        blocks.append(_classify_pdf_text(text, size, is_bold, is_mono))
                elif raw_block.get("type") == 1:
                    data = raw_block.get("image", b"")
                    if len(data) < 1024:
                        continue
                    md5 = hashlib.md5(data).hexdigest()
                    if md5 in seen_md5:
                        continue
                    ext = raw_block.get("ext", "png")
                    if ext == "jpeg":
                        ext = "jpg"
                    index = len(seen_md5)
                    seen_md5[md5] = index
                    image = ImageBlock(index=index, data=data, ext=ext, md5=md5, paragraph_idx=len(blocks))
                    images.append(image)
                    blocks.append(ContentBlock(type=BlockType.IMAGE, image=image))
    finally:
        document.close()

    text_sample = " ".join(block.content for block in blocks[:30] if block.content)
    return ParsedDocument(
        blocks=blocks,
        images=images,
        detected_lang=_detect_language(text_sample),
        title_hint=_guess_title(blocks),
    )


def blocks_to_markdown(parsed: ParsedDocument) -> str:
    parts: list[str] = []
    for block in parsed.blocks:
        if block.type == BlockType.IMAGE and block.image:
            parts.append(f"[IMG_{block.image.index}]")
        elif block.type == BlockType.HEADING:
            parts.append(f"{'#' * max(block.level, 1)} {block.content}")
        elif block.type == BlockType.CODE:
            language = block.language or ""
            parts.append(f"```{language}\n{block.content}\n```")
        elif block.content:
            parts.append(block.content)
    return "\n\n".join(part for part in parts if part)


def inject_asset_images(content: str, assets: list[dict]) -> str:
    if not assets:
        return content

    used: list[int] = []

    def image_markdown(index: int) -> str:
        if index >= len(assets):
            return ""
        asset = assets[index]
        url = asset.get("url") or asset.get("local_path") or asset.get("filename") or ""
        alt = Path(str(asset.get("filename") or f"image-{index + 1}")).stem
        return f"![{alt}]({url})"

    def replace(match: re.Match[str]) -> str:
        index = int(match.group(1))
        used.append(index)
        return image_markdown(index)

    rendered = re.sub(r"\[IMG_(\d+)\]", replace, content)
    rendered_urls = set(re.findall(r"!\[[^\]]*]\((https?://[^)]+)\)", rendered))
    missed = [index for index, _ in enumerate(assets) if index not in set(used)]
    if missed:
        fallback_images: list[str] = []
        for index in missed:
            url = assets[index].get("url") or ""
            if url and url in rendered_urls:
                continue
            fallback_images.append(image_markdown(index))
            if url:
                rendered_urls.add(url)
        if fallback_images:
            rendered = rendered.rstrip() + "\n\n" + "\n\n".join(fallback_images)
    return rendered


def normalize_published_markdown(content: str) -> str:
    content = _remove_generated_footer(content)
    content = _remove_transient_image_placeholders(content)
    content = _remove_editorial_notes(content)
    content = _remove_chunk_titles(content)
    content = _normalize_heading_levels(content)
    content = _dedupe_markdown_images(content)
    content = _repair_emphasis_markers(content)
    content = re.sub(r"\n{3,}", "\n\n", content).strip()
    return content


def render_blog_markdown(
    body: str,
    filename: str,
    assets: list[dict],
    preset: Literal["hexo", "hugo", "astro"] = "hexo",
) -> str:
    body = normalize_published_markdown(body)
    title = _first_markdown_title(body) or Path(filename).stem
    description = _description_from_body(body)
    categories, tags = _infer_taxonomy(body, filename)
    cover = next((asset.get("url") for asset in assets if asset.get("url")), "")
    now = datetime.now(timezone(timedelta(hours=8)))

    if preset == "hugo":
        date_value = now.strftime("%Y-%m-%dT%H:%M:%S+08:00")
        front_matter: dict = {
            "title": title,
            "date": date_value,
            "lastmod": date_value,
            "tags": tags,
            "categories": categories,
            "description": description,
        }
        if cover:
            front_matter["cover"] = {"image": cover}
    elif preset == "astro":
        date_value = now.strftime("%Y-%m-%dT%H:%M:%S+08:00")
        front_matter = {
            "title": title,
            "pubDate": date_value,
            "updatedDate": date_value,
            "description": description,
            "tags": tags,
            "categories": categories,
        }
        if cover:
            front_matter["image"] = cover
    else:
        date_value = now.strftime("%Y-%m-%d %H:%M:%S")
        front_matter = {
            "title": title,
            "date": date_value,
            "updated": date_value,
            "tags": tags,
            "categories": categories,
            "description": description,
        }
        if cover:
            front_matter["cover"] = cover

    yaml_text = yaml.dump(front_matter, allow_unicode=True, sort_keys=False, default_flow_style=False).rstrip()
    return f"---\n{yaml_text}\n---\n\n{body.strip()}\n"


def title_stem_from_markdown(content: str, fallback_filename: str) -> str:
    normalized = normalize_published_markdown(content)
    title = _first_markdown_title(normalized) or Path(fallback_filename).stem
    return _safe_filename_stem(title)


def _parse_docx_paragraph(
    paragraph: Paragraph,
    para_idx: int,
    seen_md5: dict[str, int],
    document: Document,
) -> tuple[ContentBlock | None, ImageBlock | None]:
    image = _extract_docx_image(paragraph, para_idx, seen_md5, document)
    if image:
        return ContentBlock(type=BlockType.IMAGE, image=image), image

    text = paragraph.text.strip()
    if not text:
        return None, None
    style_name = (paragraph.style.name if paragraph.style else "").lower()
    if style_name.startswith("heading"):
        match = re.search(r"(\d+)", style_name)
        level = min(max(int(match.group(1)) if match else 1, 1), 4)
        return ContentBlock(type=BlockType.HEADING, content=text, level=level), None
    if "code" in style_name or style_name in {"macro text", "html preformatted"}:
        return ContentBlock(type=BlockType.CODE, content=text), None
    if "list" in style_name:
        prefix = "1. " if paragraph._p.find(qn("w:numPr")) is not None else "- "
        return ContentBlock(type=BlockType.LIST, content=prefix + text), None
    return ContentBlock(type=BlockType.PARAGRAPH, content=text), None


def _extract_docx_image(
    paragraph: Paragraph,
    para_idx: int,
    seen_md5: dict[str, int],
    document: Document,
) -> ImageBlock | None:
    drawing = paragraph._p.find(".//" + qn("a:blip"))
    if drawing is None:
        return None
    embed = drawing.get(qn("r:embed"))
    if not embed:
        return None
    try:
        image_part = document.part.related_parts[embed]
    except KeyError:
        return None
    data = image_part.blob
    md5 = hashlib.md5(data).hexdigest()
    if md5 in seen_md5:
        return None
    index = len(seen_md5)
    seen_md5[md5] = index
    return ImageBlock(index=index, data=data, ext=_content_type_to_ext(image_part.content_type), md5=md5, paragraph_idx=para_idx)


def _parse_docx_table(table: Table) -> ContentBlock | None:
    if not table.rows:
        return None
    lines: list[str] = []
    for index, row in enumerate(table.rows):
        cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
        lines.append("| " + " | ".join(cells) + " |")
        if index == 0:
            lines.append("| " + " | ".join(["---"] * len(cells)) + " |")
    return ContentBlock(type=BlockType.TABLE, content="\n".join(lines))


def _extract_pdf_text_block(raw_block: dict) -> tuple[str, float, bool, bool]:
    lines: list[str] = []
    sizes: list[float] = []
    fonts: list[str] = []
    for line in raw_block.get("lines", []):
        spans = line.get("spans", [])
        text = "".join(span.get("text", "") for span in spans).strip()
        if text:
            lines.append(text)
        sizes.extend(float(span.get("size", 12)) for span in spans if span.get("size"))
        fonts.extend(str(span.get("font", "")).lower() for span in spans)
    font_text = " ".join(fonts)
    return (
        "\n".join(lines).strip(),
        sum(sizes) / len(sizes) if sizes else 12,
        any(token in font_text for token in ("bold", "black", "heavy")),
        any(token in font_text for token in ("mono", "courier", "consol", "code")),
    )


def _classify_pdf_text(text: str, size: float, is_bold: bool, is_mono: bool) -> ContentBlock:
    if is_mono:
        return ContentBlock(type=BlockType.CODE, content=text)
    if size >= 18 or (size >= 15 and is_bold):
        return ContentBlock(type=BlockType.HEADING, content=text, level=1)
    if size >= 14 or (size >= 12 and is_bold):
        return ContentBlock(type=BlockType.HEADING, content=text, level=2)
    if re.match(r"^[•·▪◦\-\*]\s+", text) or re.match(r"^\d+[.)]\s+", text):
        return ContentBlock(type=BlockType.LIST, content="- " + re.sub(r"^[•·▪◦\-\*\d.)]+\s+", "", text))
    return ContentBlock(type=BlockType.PARAGRAPH, content=text)


def _extract_docx_title(document: Document) -> str:
    for paragraph in document.paragraphs:
        style = (paragraph.style.name if paragraph.style else "").lower()
        text = paragraph.text.strip()
        if text and (style == "title" or style.startswith("heading 1")):
            return text
    return ""


def _guess_title(blocks: list[ContentBlock]) -> str:
    for block in blocks[:8]:
        if block.type == BlockType.HEADING and block.content:
            return block.content
    return ""


def _first_markdown_title(body: str) -> str:
    match = re.search(r"^#\s+(.+)$", body, re.M)
    return match.group(1).strip() if match else ""


def _safe_filename_stem(value: str) -> str:
    stem = re.sub(r"[^\w\-\u4e00-\u9fff]+", "-", value).strip("-")
    return stem[:80] or "draft-to-blog"


def _description_from_body(body: str) -> str:
    title = _first_markdown_title(body)
    plain = re.sub(r"```.*?```", "", body, flags=re.S)
    plain = re.sub(r"!\[[^\]]*]\([^)]+\)", "", plain)
    plain = re.sub(r"^\s*\[IMG_\d+]\s*$", "", plain, flags=re.M)
    lines = [_plain_description_line(line) for line in plain.splitlines()]
    title_key = _description_key(title)
    lines = [line for line in lines if line and _description_key(line) != title_key]
    topics = _extract_description_topics(lines)
    if title and topics:
        joined_topics = "、".join(topics[:4])
        return f"本文围绕《{title}》，系统梳理{joined_topics}等核心内容，帮助读者快速理解文章脉络与实践要点。"[:150]
    if title:
        return f"本文围绕《{title}》展开整理，提炼核心观点、结构脉络和实践要点，便于阅读、复盘与发布。"[:150]
    plain_text = re.sub(r"\s+", " ", " ".join(lines)).strip()
    return plain_text[:150]


def _plain_description_line(line: str) -> str:
    line = line.strip()
    if not line or re.match(r"^\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?$", line):
        return ""
    if line.startswith("|"):
        return ""
    line = re.sub(r"^#+\s*", "", line)
    line = re.sub(r"^\s*(?:[-*+]|\d+[.)、])\s+", "", line)
    line = re.sub(r"[*_`>#|\\[\]()]+", " ", line)
    return re.sub(r"\s+", " ", line).strip()


def _extract_description_topics(lines: list[str]) -> list[str]:
    topics: list[str] = []
    for line in lines:
        candidate = line.strip("：:。；;，, ")
        if not candidate:
            continue
        if "：" in candidate or ":" in candidate:
            topic = re.split(r"[:：]", candidate, maxsplit=1)[0].strip()
            if not re.search(r"[\u4e00-\u9fff]", topic):
                continue
        elif len(candidate) <= 24 and not re.search(r"[。！？.!?]", candidate):
            topic = candidate
        else:
            continue
        topic = re.sub(r"^[^\w\u4e00-\u9fff]+", "", topic).strip()
        if 2 <= len(topic) <= 24 and topic not in topics:
            topics.append(topic)
        if len(topics) >= 4:
            break
    return topics


def _description_key(value: str) -> str:
    return re.sub(r"[\W_]+", "", value).lower()


def _remove_generated_footer(content: str) -> str:
    content = re.sub(
        r"\n*\s*---\s*\n\s*>\s*本文由\s*DraftToBlog\s*完成格式整理与隐私检查，请在发布前人工复核。?\s*$",
        "",
        content,
    )
    return re.sub(r"(?m)^>\s*本文由\s*DraftToBlog\s*完成格式整理与隐私检查，请在发布前人工复核。?\s*$", "", content)


def _remove_transient_image_placeholders(content: str) -> str:
    content = re.sub(r"(?m)^@@DTBIMAGE\d+:[A-Za-z0-9_-]*@@[ \t]*$", "", content)
    content = re.sub(r"@@DTBIMAGE\d+:[A-Za-z0-9_-]*@@", "", content)
    content = re.sub(r"(?m)^draft-to-blog://[^\s)]+[ \t]*$", "", content)
    return re.sub(r"draft-to-blog://[^\s)]+", "", content)


def _remove_editorial_notes(content: str) -> str:
    blocked = (
        r"(以下是|下面是|整理结果|改写结果|注[:：]|备注[:：]|此部分原文|此部分内容|"
        r"非常好|我将|我会|作为AI|作为 AI|本文由AI|本文由 AI)"
    )
    lines = []
    for line in content.splitlines():
        stripped = line.strip()
        if re.search(blocked, stripped, re.I):
            continue
        lines.append(line)
    return "\n".join(lines)


def _remove_chunk_titles(content: str) -> str:
    lines: list[str] = []
    for line in content.splitlines():
        stripped = line.strip()
        if re.search(r"(第[一二三四五六七八九十\d]+部分|part\s+\d+)", stripped, re.I):
            cleaned = re.sub(r"[（(]?\s*第[一二三四五六七八九十\d]+部分\s*[)）]?", "", line)
            cleaned = re.sub(r"[（(]?\s*part\s+\d+\s*[)）]?", "", cleaned, flags=re.I)
            if re.match(r"^#+\s*$", cleaned.strip()):
                continue
            line = cleaned.rstrip()
        lines.append(line)
    return "\n".join(lines)


def _normalize_heading_levels(content: str) -> str:
    seen_h1 = False
    lines: list[str] = []
    for line in content.splitlines():
        if line.startswith("# "):
            if seen_h1:
                lines.append("#" + line)
            else:
                seen_h1 = True
                lines.append(line)
        else:
            lines.append(line)
    return "\n".join(lines)


def _dedupe_markdown_images(content: str) -> str:
    seen_urls: set[str] = set()
    lines: list[str] = []
    for line in content.splitlines():
        match = re.match(r"^!\[[^\]]*]\((https?://[^)]+)\)\s*$", line.strip())
        if match:
            url = match.group(1)
            if url in seen_urls:
                continue
            seen_urls.add(url)
        lines.append(line)
    return "\n".join(lines)


def _repair_emphasis_markers(content: str) -> str:
    content = re.sub(r"\*\*([^*\n]+)\*(?=[:：，,。；;\s]|$)", r"**\1**", content)
    content = re.sub(r"(?<!\*)\*([^*\n]{1,60})\*\*(?=[:：，,。；;\s]|$)", r"**\1**", content)
    return content


def _infer_taxonomy(body: str, filename: str) -> tuple[list[str], list[str]]:
    text = f"{filename} {body}".lower()
    tags: list[str] = []
    if any(word in text for word in ("数学建模", "mcm", "icm", "modeling")):
        tags.append("数学建模")
    if any(word in text for word in ("python", "matlab", "算法", "模型")):
        tags.append("技术写作")
    if not tags:
        tags.append("博客")
    return ["写作"], tags[:5]


def _detect_language(text: str) -> str:
    cjk = sum(1 for char in text if "\u4e00" <= char <= "\u9fff")
    return "zh" if cjk / max(len(text), 1) > 0.1 else "en"


def _content_type_to_ext(content_type: str) -> str:
    return {
        "image/png": "png",
        "image/jpeg": "jpg",
        "image/gif": "gif",
        "image/webp": "webp",
        "image/bmp": "bmp",
    }.get(content_type.lower(), "png")
