from __future__ import annotations

import re
import uuid
from io import BytesIO
from pathlib import Path
from typing import Any

import fitz
from docx import Document

from ..storage import store
from .ai_service import improve_article, resolve_connection
from .media_service import append_uploaded_assets, upload_task_assets


SUPPORTED_EXTENSIONS = {".md": "markdown", ".markdown": "markdown", ".docx": "docx", ".pdf": "pdf"}
IMAGE_PATTERN = re.compile(r"!\[[^\]]*\]\(([^)]+)\)")

SENSITIVE_PATTERNS: tuple[tuple[str, str, str, str, re.Pattern[str]], ...] = (
    ("phone", "手机号", "high", "138****5678", re.compile(r"(?<!\d)1[3-9]\d{9}(?!\d)")),
    ("email", "邮箱", "medium", "u***@example.com", re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")),
    ("id_card", "身份证号", "high", "[身份证号]", re.compile(r"(?<!\d)\d{17}[\dXx](?!\d)")),
    ("ipv4", "IP 地址", "high", "[内网 IP]", re.compile(r"(?<!\d)(?:\d{1,3}\.){3}\d{1,3}(?!\d)")),
    ("internal_url", "内网地址", "high", "[内网地址]", re.compile(r"https?://(?:localhost|[^\s/]*(?:\.local|\.internal)|10\.\d+\.\d+\.\d+|192\.168\.\d+\.\d+)[^\s)]*", re.I)),
)


def analyze_upload(filename: str, content: bytes) -> dict[str, Any]:
    extension = Path(filename).suffix.lower()
    source_type = SUPPORTED_EXTENSIONS.get(extension)
    if not source_type:
        raise ValueError("首版仅支持 PDF、DOCX 和 Markdown 文件。")

    task_id = uuid.uuid4().hex
    warnings: list[str] = []
    if source_type == "markdown":
        text, image_count, markdown_warnings = _parse_markdown(content)
        warnings.extend(markdown_warnings)
    elif source_type == "docx":
        text, images = _parse_docx(content)
        image_count = len(images)
    else:
        text, images, pdf_warnings = _parse_pdf(content)
        image_count = len(images)
        warnings.extend(pdf_warnings)

    if not text.strip():
        raise ValueError("没有从文件中提取到可处理的文字内容。")

    findings = find_sensitive_content(text)
    asset_records = [
        {
            "filename": f"image-{index}{suffix}",
            "local_path": f"assets/image-{index}{suffix}",
            "status": "pending",
            "url": "",
            "provider": "",
            "error": "",
        }
        for index, (suffix, _) in enumerate(images if source_type in {"docx", "pdf"} else [], start=1)
    ]
    payload: dict[str, Any] = {
        "id": task_id,
        "filename": Path(filename).name,
        "source_type": source_type,
        "status": "ready",
        "original_content": text,
        "processed_content": text,
        "stats": {
            "characters": len(text),
            "paragraphs": len([line for line in text.splitlines() if line.strip()]),
            "images": image_count,
            "findings": len(findings),
        },
        "findings": findings,
        "warnings": warnings,
        "assets": asset_records,
    }
    store.create(task_id, payload)

    if source_type in {"docx", "pdf"}:
        for index, (suffix, image_bytes) in enumerate(images, start=1):
            store.save_asset(task_id, f"image-{index}{suffix}", image_bytes)
    return payload


def create_demo() -> dict[str, Any]:
    content = """# Apollo 项目故障复盘

昨天线上告警以后，我们先登录内部面板查看。负责人张三的联系方式是 13812345678，通知邮箱是 zhangsan@example.com。

## 排查记录

服务地址为 http://10.20.30.40:8080/admin。最开始怀疑缓存，后来确认是配置发布顺序错误。

## 处理结果

重新发布配置后服务恢复。后续需要整理成可以公开分享的排查教程。
"""
    payload = analyze_upload("DraftToBlog-演示文档.md", content.encode("utf-8"))
    payload["warnings"].append("演示文档使用虚构信息，可安全体验脱敏流程。")
    store.update(payload["id"], payload)
    return payload


def find_sensitive_content(text: str) -> list[dict[str, Any]]:
    findings: list[dict[str, Any]] = []
    occupied: list[tuple[int, int]] = []
    for kind, label, risk, replacement, pattern in SENSITIVE_PATTERNS:
        for match in pattern.finditer(text):
            if any(match.start() < end and match.end() > start for start, end in occupied):
                continue
            original = match.group(0)
            masked = replacement
            if kind == "phone":
                masked = f"{original[:3]}****{original[-4:]}"
            elif kind == "email":
                local, domain = original.split("@", 1)
                masked = f"{local[:1]}***@{domain}"
            findings.append(
                {
                    "id": uuid.uuid4().hex,
                    "kind": kind,
                    "label": label,
                    "original": original,
                    "replacement": masked,
                    "start": match.start(),
                    "end": match.end(),
                    "risk": risk,
                    "accepted": True,
                }
            )
            occupied.append((match.start(), match.end()))
    return sorted(findings, key=lambda item: item["start"])


def process_document(
    task_id: str,
    finding_ids: list[str],
    improve_structure: bool,
    ai_config: object | None = None,
) -> dict[str, Any]:
    payload = store.get(task_id)
    if not payload:
        raise FileNotFoundError(task_id)

    selected = {item["id"] for item in payload["findings"] if item["id"] in finding_ids}
    text = payload["original_content"]
    replacements = [item for item in payload["findings"] if item["id"] in selected]
    for finding in sorted(replacements, key=lambda item: item["start"], reverse=True):
        text = text[: finding["start"]] + finding["replacement"] + text[finding["end"] :]

    assets = payload.get("assets", [])
    if assets:
        upload_task_assets(task_id, assets)
        payload["assets"] = assets
        # Persist successful uploads before calling AI so retries never upload
        # the same document images again when the model fails or times out.
        store.update(task_id, payload)
        failed = [item for item in assets if item.get("status") == "failed"]
        if failed:
            message = f"{len(failed)} 张图片上传失败，已保留本地任务文件。"
            if message not in payload["warnings"]:
                payload["warnings"].append(message)

    if improve_structure:
        connection = resolve_connection(ai_config)
        if connection:
            text = improve_article(text, payload["filename"], connection)
        else:
            text = _improve_structure(text, payload["filename"])

    if assets:
        text = append_uploaded_assets(text, assets)

    for finding in payload["findings"]:
        finding["accepted"] = finding["id"] in selected
    payload["processed_content"] = text
    payload["status"] = "processed"
    store.update(task_id, payload)
    return payload


def _parse_markdown(content: bytes) -> tuple[str, int, list[str]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError("Markdown 需要使用 UTF-8 编码。") from exc
    refs = IMAGE_PATTERN.findall(text)
    warnings: list[str] = []
    local_refs = [ref for ref in refs if not re.match(r"^(?:https?://|data:)", ref, re.I)]
    if not refs:
        warnings.append("未检测到图片。飞书导出的 Markdown 可能没有携带图片资源。")
    elif local_refs:
        warnings.append(f"检测到 {len(local_refs)} 个本地图片引用；上传 COS 前需要补充对应文件。")
    return text, len(refs), warnings


def _parse_docx(content: bytes) -> tuple[str, list[tuple[str, bytes]]]:
    try:
        document = Document(BytesIO(content))
    except Exception as exc:
        raise ValueError("DOCX 文件无法读取，可能已损坏或加密。") from exc

    blocks: list[str] = []
    for paragraph in document.paragraphs:
        value = paragraph.text.strip()
        if not value:
            continue
        style = paragraph.style.name.lower() if paragraph.style else ""
        if style.startswith("heading"):
            level_match = re.search(r"(\d+)", style)
            level = min(int(level_match.group(1)), 6) if level_match else 2
            blocks.append(f"{'#' * level} {value}")
        else:
            blocks.append(value)
    for table in document.tables:
        for row in table.rows:
            blocks.append(" | ".join(cell.text.strip() for cell in row.cells))

    images: list[tuple[str, bytes]] = []
    for relationship in document.part.rels.values():
        if "image" not in relationship.reltype:
            continue
        part = relationship.target_part
        suffix = Path(str(part.partname)).suffix or ".png"
        images.append((suffix, part.blob))
    return "\n\n".join(blocks), images


def _parse_pdf(content: bytes) -> tuple[str, list[tuple[str, bytes]], list[str]]:
    try:
        document = fitz.open(stream=content, filetype="pdf")
    except Exception as exc:
        raise ValueError("PDF 文件无法读取，可能已损坏或加密。") from exc
    if document.needs_pass:
        raise ValueError("首版暂不支持加密 PDF，请先解除密码。")

    pages: list[str] = []
    images: list[tuple[str, bytes]] = []
    seen_xrefs: set[int] = set()
    for page in document:
        pages.append(page.get_text("text").strip())
        for image in page.get_images(full=True):
            xref = image[0]
            if xref in seen_xrefs:
                continue
            seen_xrefs.add(xref)
            extracted = document.extract_image(xref)
            images.append((f".{extracted.get('ext', 'png')}", extracted["image"]))
    text = "\n\n".join(value for value in pages if value)
    warnings: list[str] = []
    if not text.strip():
        warnings.append("该 PDF 可能是扫描件；首版暂未启用 OCR。")
    return text, images, warnings


def _improve_structure(text: str, filename: str) -> str:
    normalized = re.sub(r"\n{3,}", "\n\n", text.strip())
    if normalized.startswith("# "):
        return normalized + "\n\n---\n\n> 本文由 DraftToBlog 完成格式整理与隐私检查，请在发布前人工复核。"
    title = Path(filename).stem.replace("-", " ").strip() or "整理后的文章"
    return f"# {title}\n\n{normalized}\n\n---\n\n> 本文由 DraftToBlog 完成格式整理与隐私检查，请在发布前人工复核。"
